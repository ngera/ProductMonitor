"""Convert a Markdown file to a Word (.docx) document.

Usage:
    python scripts/md_to_docx.py <input.md> <output.docx>

Supports the markdown subset used by this project's design docs:
    - # / ## / ### / #### / ##### headings
    - Paragraphs
    - Bullet lists ("- " or "* ")
    - Numbered lists ("1. ")
    - Tables with "| ... |" rows and "| --- |" separator
    - Fenced code blocks (``` ... ```)
    - Horizontal rules (--- on its own line)
    - Blockquotes (lines starting with "> ")
    - Inline **bold**, *italic*, `code`, and [text](url) links

Anything outside this subset is rendered as a plain paragraph with its
markdown source intact, so the output is not lossy in a visible way.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

_HEADING_RE = re.compile(r"^(#{1,5})\s+(.*)$")
_FENCE_RE = re.compile(r"^```")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")
_TABLE_ROW_RE = re.compile(r"^\s*\|.+\|\s*$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\d+\.\s+(.*)$")
_BLOCKQUOTE_RE = re.compile(r"^>\s?(.*)$")
_HR_RE = re.compile(r"^-{3,}\s*$")

# Inline tokenizer: keeps **bold**, *italic*, `code`, [text](url).
_INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)


def _add_runs_with_formatting(paragraph, text: str) -> None:
    """Append text to `paragraph` honoring inline markdown formatting."""
    parts = _INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            text_end = part.index("](")
            link_text = part[1:text_end]
            url = part[text_end + 2 : -1]
            run = paragraph.add_run(link_text)
            run.font.color.rgb = RGBColor(0x0B, 0x5C, 0xAD)
            run.font.underline = True
            # Hyperlink relationship is best-effort here; the visible
            # text + URL coloring is enough for a spec. Embedding a
            # real hyperlink relationship requires more OOXML plumbing.
        else:
            paragraph.add_run(part)


def _add_paragraph(doc: Document, text: str, *, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    _add_runs_with_formatting(p, text)


def _add_code_block(doc: Document, lines: list[str]) -> None:
    body = "\n".join(lines)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(14)
    run = p.add_run(body)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    # Light gray shading for the paragraph.
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F4F4F4")
    pPr.append(shd)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Light Grid Accent 1"
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.rows[r_idx].cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            cell_text = row[c_idx] if c_idx < len(row) else ""
            cell.text = ""  # clear default paragraph
            para = cell.paragraphs[0]
            _add_runs_with_formatting(para, cell_text)
            if r_idx == 0:
                for run in para.runs:
                    run.bold = True


def _split_table_row(line: str) -> list[str]:
    # Trim leading/trailing pipe, split, strip cells.
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [c.strip() for c in stripped.split("|")]


def _add_hr(doc: Document) -> None:
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    pBdr.append(bottom)
    pPr.append(pBdr)


def convert(md_path: Path, out_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    # Set base font.
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]

        # Fenced code block
        if _FENCE_RE.match(line):
            i += 1
            block: list[str] = []
            while i < len(lines) and not _FENCE_RE.match(lines[i]):
                block.append(lines[i])
                i += 1
            _add_code_block(doc, block)
            i += 1  # skip closing ```
            continue

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if _HR_RE.match(line):
            _add_hr(doc)
            i += 1
            continue

        # Heading
        m = _HEADING_RE.match(line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            doc.add_heading(text, level=level)
            i += 1
            continue

        # Table — detect header row followed by separator row
        if _TABLE_ROW_RE.match(line) and i + 1 < len(lines) and _TABLE_SEP_RE.match(lines[i + 1]):
            rows: list[list[str]] = [_split_table_row(line)]
            i += 2  # skip header and separator
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i]):
                rows.append(_split_table_row(lines[i]))
                i += 1
            _add_table(doc, rows)
            continue

        # Blockquote
        bq_match = _BLOCKQUOTE_RE.match(line)
        if bq_match:
            quoted_lines = [bq_match.group(1)]
            i += 1
            while i < len(lines):
                m2 = _BLOCKQUOTE_RE.match(lines[i])
                if not m2:
                    break
                quoted_lines.append(m2.group(1))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            run = p.add_run(" ".join(quoted_lines))
            run.italic = True
            continue

        # Bullet list (consume consecutive bullet lines)
        if _BULLET_RE.match(line):
            while i < len(lines):
                m2 = _BULLET_RE.match(lines[i])
                if not m2:
                    break
                _add_paragraph(doc, m2.group(1), style="List Bullet")
                i += 1
            continue

        # Numbered list
        if _NUMBERED_RE.match(line):
            while i < len(lines):
                m2 = _NUMBERED_RE.match(lines[i])
                if not m2:
                    break
                _add_paragraph(doc, m2.group(1), style="List Number")
                i += 1
            continue

        # Default: paragraph (may span multiple consecutive non-blank lines)
        paragraph_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not (
            _HEADING_RE.match(lines[i])
            or _FENCE_RE.match(lines[i])
            or _HR_RE.match(lines[i])
            or _BULLET_RE.match(lines[i])
            or _NUMBERED_RE.match(lines[i])
            or _BLOCKQUOTE_RE.match(lines[i])
            or (_TABLE_ROW_RE.match(lines[i]) and i + 1 < len(lines)
                and _TABLE_SEP_RE.match(lines[i + 1]))
        ):
            paragraph_lines.append(lines[i])
            i += 1
        _add_paragraph(doc, " ".join(paragraph_lines))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: python scripts/md_to_docx.py <input.md> <output.docx>",
              file=sys.stderr)
        return 2
    md_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])
    if not md_path.exists():
        print(f"input file not found: {md_path}", file=sys.stderr)
        return 1
    convert(md_path, out_path)
    print(f"wrote {out_path} ({out_path.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
